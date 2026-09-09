import csv
import io
from datetime import date, timedelta

from flask import Blueprint, render_template, request, send_file

from app import db, pdf_reports, branding
from app.auth import login_required

bp = Blueprint("reports", __name__, url_prefix="/reports")


def _logo_path():
    return branding.get_logo_path()


def _default_range():
    today = date.today()
    start = today.replace(day=1)
    return start.isoformat(), today.isoformat()


def _resolve_range():
    start = request.args.get("start", "").strip()
    end   = request.args.get("end",   "").strip()
    if not start or not end:
        start, end = _default_range()
    return start, end


def _compute_stats(start, end):
    return {
        "revenue":      db.get_revenue_collected(start, end),
        "cases_closed": db.get_cases_closed_count(start, end),
        "new_cases":    db.get_new_cases_count(start, end),
        "new_patients": db.get_new_patients_count(start, end),
    }


@bp.route("/")
@login_required
def view_reports():
    start, end = _resolve_range()
    stats          = _compute_stats(start, end)
    payments       = db.get_payments_in_range(start, end)
    closed_cases   = db.get_cases_closed_in_range(start, end)
    monthly_series = db.get_monthly_revenue_series(date.today().year)
    revenue_summary = db.get_revenue_summary()
    return render_template(
        "reports.html", start=start, end=end, stats=stats, payments=payments,
        closed_cases=closed_cases, monthly_series=monthly_series,
        current_year=date.today().year, revenue_summary=revenue_summary,
        today_str=date.today().isoformat(),
    )


@bp.route("/pdf")
@login_required
def export_pdf():
    start, end   = _resolve_range()
    stats        = _compute_stats(start, end)
    payments     = db.get_payments_in_range(start, end)
    closed_cases = db.get_cases_closed_in_range(start, end)
    clinic       = db.get_all_settings()
    revenue_summary = db.get_revenue_summary()
    pdf_bytes = pdf_reports.generate_report_pdf(
        start, end, stats, payments, closed_cases, clinic, _logo_path(),
        revenue_summary=revenue_summary,
    )
    filename = f"feast9_report_{start}_to_{end}.pdf"
    return send_file(io.BytesIO(pdf_bytes), mimetype="application/pdf",
                      as_attachment=False, download_name=filename)


@bp.route("/csv")
@login_required
def export_csv():
    """Payments received in the selected date range."""
    start, end = _resolve_range()
    payments   = db.get_payments_in_range(start, end)
    buf        = io.StringIO()
    writer     = csv.writer(buf)
    writer.writerow(["Date", "Patient", "Case", "Amount (Rs.)", "Method", "Note"])
    for p in payments:
        writer.writerow([p["payment_date"], p.get("patient_name",""),
                          p.get("case_title",""), p["amount"],
                          p.get("method") or "", p.get("note") or ""])
    mem      = io.BytesIO(buf.getvalue().encode("utf-8"))
    filename = f"feast9_payments_{start}_to_{end}.csv"
    return send_file(mem, mimetype="text/csv", as_attachment=True, download_name=filename)


@bp.route("/pending-csv")
@login_required
def export_pending_csv():
    """All cases with outstanding balances — mirrors the pending payments view."""
    revenue_summary = db.get_revenue_summary()
    pending = revenue_summary.get("patient_case_pending", [])
    today   = date.today().isoformat()
    buf     = io.StringIO()
    writer  = csv.writer(buf)
    writer.writerow(["Patient", "Mobile", "Case", "Doctor", "Status",
                      "Total Billed (Rs.)", "Received (Rs.)", "Pending (Rs.)",
                      "Follow-up Date", "Overdue?"])
    for r in pending:
        overdue = "Yes" if (r.get("follow_up_date") and r["follow_up_date"] < today) else "No"
        writer.writerow([r["patient_name"], r["mobile"], r["case_title"],
                          r.get("doctor_name") or "", r["status"],
                          r["total_cost"], r["paid"], r["pending"],
                          r.get("follow_up_date") or "", overdue])
    # Totals row
    writer.writerow(["TOTAL", "", "", "", "",
                      sum(r["total_cost"] for r in pending),
                      sum(r["paid"] for r in pending),
                      sum(r["pending"] for r in pending), "", ""])
    mem      = io.BytesIO(buf.getvalue().encode("utf-8"))
    filename = f"feast9_pending_payments_{date.today().isoformat()}.csv"
    return send_file(mem, mimetype="text/csv", as_attachment=True, download_name=filename)


@bp.route("/pending-excel")
@login_required
def export_pending_excel():
    """Excel workbook with two sheets — Pending Payments and Revenue Summary."""
    revenue_summary = db.get_revenue_summary()
    pending = revenue_summary.get("patient_case_pending", [])
    today   = date.today().isoformat()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    GREEN_DARK  = "2E7D32"
    GREEN_LIGHT = "E8F5E9"
    RED_LIGHT   = "FDECEA"
    AMBER_LIGHT = "FFF8E1"
    WHITE       = "FFFFFF"

    wb = Workbook()

    def hdr(cell, bg=GREEN_DARK):
        cell.font      = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cell.fill      = PatternFill("solid", start_color=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def thin():
        s = Side(style="thin", color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)

    def cell_val(cell, val, bg=None, bold=False, align="left", fmt=None):
        cell.value     = val if val is not None else ""
        cell.font      = Font(name="Calibri", size=9, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border    = thin()
        if bg:  cell.fill        = PatternFill("solid", start_color=bg)
        if fmt: cell.number_format = fmt

    # ── Sheet 1: Pending Payments ─────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Pending Payments"
    ws1.freeze_panes = "A2"

    HEADERS = [
        ("Patient Name", 22), ("Mobile", 14), ("Case", 26), ("Doctor", 18),
        ("Status", 10), ("Total Billed (Rs.)", 16), ("Received (Rs.)", 15),
        ("Pending (Rs.)", 15), ("Follow-up Date", 14), ("Overdue?", 10),
    ]
    ws1.row_dimensions[1].height = 32
    for ci, (label, width) in enumerate(HEADERS, 1):
        hdr(ws1.cell(row=1, column=ci, value=label))
        ws1.column_dimensions[get_column_letter(ci)].width = width

    for ri, r in enumerate(pending, 2):
        overdue = bool(r.get("follow_up_date") and r["follow_up_date"] < today)
        bg = RED_LIGHT if overdue else (
            AMBER_LIGHT if r.get("status") == "Active" else WHITE)
        vals = [
            r["patient_name"], r["mobile"], r["case_title"],
            r.get("doctor_name") or "", r["status"],
            r["total_cost"], r["paid"], r["pending"],
            r.get("follow_up_date") or "", "YES — OVERDUE" if overdue else "",
        ]
        for ci, val in enumerate(vals, 1):
            c = ws1.cell(row=ri, column=ci)
            is_money = ci in (6, 7, 8)
            cell_val(c, val, bg=bg, align="right" if is_money else "left",
                      fmt="#,##0.00" if is_money else None)
        ws1.row_dimensions[ri].height = 16

    # Totals row
    tr = len(pending) + 2
    totals = [("TOTAL", None, None, None, None,
                sum(r["total_cost"] for r in pending),
                sum(r["paid"]       for r in pending),
                sum(r["pending"]    for r in pending),
                None, None)]
    for ci, val in enumerate(totals[0], 1):
        c = ws1.cell(row=tr, column=ci)
        is_money = ci in (6, 7, 8)
        cell_val(c, val, bg=GREEN_LIGHT, bold=True,
                  align="right" if is_money else "left",
                  fmt="#,##0.00" if is_money else None)

    # ── Sheet 2: Revenue Summary ──────────────────────────────────────────
    ws2 = wb.create_sheet("Revenue Summary")
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 22

    summary_rows = [
        ("Metric", "Value"),
        ("Total Billed (all cases)",     revenue_summary["total_billed"]),
        ("Total Collected",               revenue_summary["total_collected"]),
        ("Total Pending",                 revenue_summary["total_pending"]),
        ("Collection Rate (%)",           f"{revenue_summary['collection_rate']}%"),
        ("Pending on Active Cases",       revenue_summary["pending_active"]),
        ("Active Cases with Pending",     revenue_summary["active_case_count"]),
        ("Overdue Cases with Balance",    revenue_summary["overdue_outstanding"]),
        ("As of",                         today),
    ]
    for ri, (label, val) in enumerate(summary_rows, 1):
        ca = ws2.cell(row=ri, column=1, value=label)
        cb = ws2.cell(row=ri, column=2, value=val)
        if ri == 1:
            hdr(ca); hdr(cb)
        else:
            is_money = isinstance(val, float) and ri not in (4,)
            for c in (ca, cb):
                c.font   = Font(name="Calibri", size=9, bold=(ri == 4 or ri == 3))
                c.border = thin()
                c.alignment = Alignment(horizontal="left" if c == ca else "right",
                                         vertical="center")
                c.fill = PatternFill("solid", start_color=(
                    RED_LIGHT if label == "Total Pending" else
                    GREEN_LIGHT if label == "Total Collected" else WHITE))
            if is_money:
                cb.number_format = "#,##0.00"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"feast9_pending_report_{date.today().isoformat()}.xlsx"
    return send_file(buf,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                      as_attachment=True, download_name=filename)
