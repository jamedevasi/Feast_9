"""
One-time historic patient data importer.

Reads an uploaded .xlsx file, maps its header row onto our patient fields
(tolerant of common header spellings/synonyms), validates every row, and
bulk-inserts the valid ones. Invalid rows are skipped and reported back to
the practitioner with the exact reason, rather than blocking the whole
import — this matters at ~8000-row scale where a handful of dirty rows are
expected.
"""
import io
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app import db
from app.validators import (
    clean_str, is_valid_mobile, is_valid_email, is_valid_age, normalize_date,
)

# Canonical field -> accepted header spellings (normalized: lowercase, only a-z0-9)
HEADER_SYNONYMS = {
    "name": {"name", "patientname", "fullname", "patientfullname"},
    "age": {"age"},
    "sex": {"sex", "gender"},
    "mobile": {"mobile", "mobilenumber", "phone", "phonenumber", "contact",
               "contactnumber", "contactno", "mobileno"},
    "email": {"email", "emailaddress", "email id", "emailid"},
    "address": {"address", "residentialaddress", "homeaddress"},
    "last_visited_date": {"lastvisiteddate", "lastvisit", "lastvisitdate",
                           "lastvisited", "datelastvisited"},
    # Optional — only present if the historic sheet happened to track these.
    # Free text only; doesn't try to match against the standardized
    # checklists, since historic notes rarely line up with a fixed list.
    "medical_conditions_other": {"medicalhistory", "medicalconditions",
                                  "medicalcondition", "healthhistory"},
    "allergies_other": {"allergies", "allergy", "drugallergies", "knownallergies"},
    "emergency_contact_name": {"emergencycontactname", "emergencycontact",
                                "emergencyname"},
    "emergency_contact_number": {"emergencycontactnumber", "emergencycontactno",
                                  "emergencyphone", "emergencynumber"},
}

REQUIRED_FIELDS = ("name", "mobile")
TEMPLATE_HEADERS = ["Name", "Age", "Sex", "Mobile", "Email", "Address", "Last Visited Date"]


def _normalize_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _map_headers(header_row) -> dict:
    """Returns {column_index: canonical_field_name} for recognized columns."""
    norm_to_field = {}
    for field, synonyms in HEADER_SYNONYMS.items():
        for syn in synonyms:
            norm_to_field[_normalize_header(syn)] = field

    mapping = {}
    for idx, cell in enumerate(header_row):
        norm = _normalize_header(cell)
        if norm in norm_to_field:
            mapping[idx] = norm_to_field[norm]
    return mapping


def preview_import(file_bytes: bytes, max_preview_rows: int = 20) -> dict:
    """Parse the file and validate every row WITHOUT writing to the database.
    Returns a summary the UI can show before the practitioner confirms."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return {"error": "The file appears to be empty."}

    mapping = _map_headers(header_row)
    missing = [f for f in REQUIRED_FIELDS if f not in mapping.values()]
    if missing:
        return {
            "error": (
                f"Could not find required column(s): {', '.join(missing)}. "
                f"Detected columns: {', '.join(str(h) for h in header_row if h)}. "
                "Please use the provided template or make sure 'Name' and "
                "'Mobile' columns are present."
            )
        }

    valid_rows, error_rows = [], []
    total = 0
    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        total += 1
        record = {f: None for f in HEADER_SYNONYMS}
        for idx, field in mapping.items():
            record[field] = row[idx] if idx < len(row) else None

        name = clean_str(record.get("name"))
        mobile = clean_str(record.get("mobile"))
        sex = clean_str(record.get("sex")).title()
        if sex not in ("Male", "Female", "Other"):
            sex = sex or ""
        email = clean_str(record.get("email"))
        address = clean_str(record.get("address"))
        age_raw = record.get("age")

        errs = []
        if not name:
            errs.append("Missing name")
        if not mobile:
            errs.append("Missing mobile number")
        elif not is_valid_mobile(mobile):
            errs.append("Invalid mobile number")
        if email and not is_valid_email(email):
            errs.append("Invalid email")
        if age_raw not in (None, "") and not is_valid_age(age_raw):
            errs.append("Invalid age")
        try:
            last_visit = normalize_date(record.get("last_visited_date"))
        except ValueError:
            errs.append("Unrecognized last-visited date format")
            last_visit = ""

        clean_record = {
            "name": name, "age": int(age_raw) if str(age_raw or "").strip().isdigit() else None,
            "sex": sex, "mobile": mobile, "email": email, "address": address,
            "last_visited_date": last_visit,
            "medical_conditions_other": clean_str(record.get("medical_conditions_other")),
            "allergies_other": clean_str(record.get("allergies_other")),
            "emergency_contact_name": clean_str(record.get("emergency_contact_name")),
            "emergency_contact_number": clean_str(record.get("emergency_contact_number")),
        }
        if errs:
            error_rows.append({"row": row_num, "errors": errs, "data": clean_record})
        else:
            valid_rows.append(clean_record)

    return {
        "total_rows": total,
        "valid_count": len(valid_rows),
        "error_count": len(error_rows),
        "valid_rows": valid_rows,
        "error_rows": error_rows[:200],  # cap what we render; full count still shown
        "preview": valid_rows[:max_preview_rows],
    }


def commit_import(valid_rows: list) -> int:
    """Insert previously-validated rows in a single transaction. Returns
    count inserted."""
    return db.bulk_add_patients(valid_rows)


def build_template_xlsx() -> bytes:
    """Generates the downloadable .xlsx template practitioners should fill in."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Patients"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="43A047")
        cell.alignment = Alignment(horizontal="center")
    ws.append(["Ravi Kumar", 34, "Male", "9876543210", "ravi@example.com", "Kochi, Kerala", "2026-05-12"])
    widths = [22, 6, 10, 16, 26, 30, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
